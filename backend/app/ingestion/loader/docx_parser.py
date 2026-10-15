"""
Word (.docx) 解析器。
"""
from pathlib import Path

from docx import Document as DocxDocument

from app.core.schemas import Document
from app.ingestion.loader.base import BaseParser


class DocxParser(BaseParser):
    SUPPORTED_EXT = {".docx"}

    def parse(self, file_path: Path) -> list[Document]:
        doc = DocxDocument(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        return [
            Document(
                filename=file_path.name,
                text=text,
                metadata={"paragraph_count": len(paragraphs)},
            )
        ]
